"""Inventory import from Excel (.xlsx) or CSV.

Flow: upload -> parse -> validate every row -> preview with row-wise errors ->
correct rows -> re-validate -> commit (all-or-nothing, inside one transaction)."""
import csv
import io
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import Category, ImportBatch, Product, ProductMappingHistory, StockMovement, Supplier, User, now
from .services import BizError, audit, norm_catalogue, norm_no
from .types import d

FIELDS = ["catalogue", "catalogue_no", "factory_order_no", "name", "brand", "hsn", "category", "description", "size", "finish",
          "unit", "sale_price", "purchase_cost", "tax_rate", "quantity", "reorder_level", "supplier"]
REQUIRED = ["catalogue_no", "factory_order_no", "name"]
NUMERIC = {"sale_price": "0.01", "purchase_cost": "0.01", "tax_rate": "0.01", "quantity": "0.001",
           "reorder_level": "0.001"}

ALIASES = {
    "catalogue": ["catalogue", "catalog", "catalogue_name", "catalog_name", "book", "catalogue_book"],
    "catalogue_no": ["catalogue_no", "catalog_no", "customer_catalogue_number", "customer_catalog_number", "design",
                     "design_code", "design_/_code", "design/code", "code",
                     "catalogue_number", "catalog_number", "customer_no", "customer_number", "cat_no", "design_no"],
    "factory_order_no": ["factory_order_no", "factory_order_number", "factory_no", "factory_number",
                         "order_no", "manufacturer_no", "factory_code"],
    "name": ["name", "product_name", "product", "item", "item_name"],
    "brand": ["brand", "company", "make", "brand_name"],
    "hsn": ["hsn", "hsn_code", "hsn_sac"],
    "category": ["category", "type", "product_type"],
    "description": ["description", "details", "desc"],
    "size": ["size", "dimensions", "dimension"],
    "finish": ["finish", "texture"],
    "unit": ["unit", "uom"],
    "sale_price": ["sale_price", "price", "rate", "selling_price", "mrp"],
    "purchase_cost": ["purchase_cost", "cost", "cost_price", "purchase_price"],
    "tax_rate": ["tax_rate", "gst", "gst_rate", "tax", "gst_%", "tax_%"],
    "quantity": ["quantity", "qty", "available_quantity", "stock", "available_qty", "opening_stock"],
    "reorder_level": ["reorder_level", "reorder", "min_stock", "minimum_stock"],
    "supplier": ["supplier", "factory", "manufacturer", "supplier_name"],
}
_lookup = {alias: field for field, aliases in ALIASES.items() for alias in aliases}


def _key(header) -> str:
    return str(header or "").strip().lower().replace(".", "").replace("-", "_").replace(" ", "_")


def parse_file(filename: str, content: bytes) -> tuple[list[dict], list[str]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = next(it, None) or []
        raw_rows = list(it)
        wb.close()
    elif name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        headers = next(reader, None) or []
        raw_rows = list(reader)
    else:
        raise BizError("Upload an Excel (.xlsx) or CSV file.")
    mapping, unknown = {}, []
    for i, h in enumerate(headers):
        k = _key(h)
        if not k:
            continue
        f = _lookup.get(k)
        if f and f not in mapping.values():
            mapping[i] = f
        else:
            unknown.append(str(h))
    missing = [f for f in REQUIRED if f not in mapping.values()]
    if missing:
        raise BizError("These columns are missing: " + ", ".join(missing) +
                       ". Download the template to see the expected column names.")
    rows = []
    for n, raw in enumerate(raw_rows, start=2):
        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
            continue
        data = {f: "" for f in FIELDS}
        for i, f in mapping.items():
            v = raw[i] if i < len(raw) else None
            if v is None:
                v = ""
            elif isinstance(v, float) and v.is_integer() and f in ("catalogue_no", "factory_order_no"):
                v = str(int(v))
            data[f] = str(v).strip()
        rows.append({"row": n, "data": data, "errors": [], "warnings": [], "action": ""})
    if not rows:
        raise BizError("The file has no product rows.")
    return rows, unknown


def validate_rows(db: Session, rows: list[dict], mode: str) -> dict:
    """Validate all rows together (needed to catch duplicates inside the file)."""
    existing_by_cat = {}
    existing_by_fac = {}
    for pid, cat, cno, fno, sid, stock in db.execute(select(Product.id, Product.catalogue, Product.catalogue_no,
                                                            Product.factory_order_no, Product.supplier_id,
                                                            Product.stock_qty)):
        existing_by_cat[(cat, cno)] = (pid, fno, sid, stock)
        existing_by_fac[(sid, fno)] = (pid, cno, cat)
    suppliers = {s.name.lower(): s.id for s in db.execute(select(Supplier)).scalars()}

    seen_cat: dict[tuple, list] = {}
    seen_fac: dict[tuple, list] = {}
    for r in rows:
        data, errs, warns = r["data"], [], []
        data["catalogue"] = norm_catalogue(data.get("catalogue"))
        data["catalogue_no"] = norm_no(data.get("catalogue_no"))
        data["factory_order_no"] = norm_no(data.get("factory_order_no"))
        data["name"] = str(data.get("name") or "").strip()
        data["supplier"] = str(data.get("supplier") or "").strip() or "Default Factory"
        if not data["catalogue_no"]:
            errs.append("Customer catalogue number is missing.")
        if not data["factory_order_no"]:
            errs.append("Factory order number is missing (no mapping).")
        if not data["name"]:
            errs.append("Product name is missing.")
        if len(data["catalogue_no"]) > 60 or len(data["factory_order_no"]) > 60:
            errs.append("Product numbers must be 60 characters or fewer.")
        for f, places in NUMERIC.items():
            v = str(data.get(f) or "").strip().replace(",", "")
            if v == "":
                continue
            try:
                num = Decimal(v)
                if num < 0:
                    errs.append(f"{f.replace('_', ' ').capitalize()} cannot be negative.")
                elif f == "tax_rate" and num > 100:
                    errs.append("Tax rate must be between 0 and 100.")
                data[f] = str(d(num, places))
            except InvalidOperation:
                errs.append(f"{f.replace('_', ' ').capitalize()} '{v}' is not a number.")
        sid = suppliers.get(data["supplier"].lower())  # None = new supplier, created on commit
        ckey = (data["catalogue"], data["catalogue_no"])
        fkey = (data["supplier"].lower(), data["factory_order_no"])
        if data["catalogue_no"]:
            seen_cat.setdefault(ckey, []).append(r)
        if data["factory_order_no"]:
            seen_fac.setdefault(fkey, []).append(r)

        existing = existing_by_cat.get(ckey) if data["catalogue_no"] else None
        if existing:
            if mode == "create":
                errs.append(f"Catalogue number {data['catalogue_no']} already exists. Use 'Add new and update existing' to update it.")
            r["action"] = "UPDATE"
            r["product_id"] = existing[0]
            if data["factory_order_no"] and (existing[1] != data["factory_order_no"] or (sid and existing[2] != sid)):
                warns.append(f"Mapping will change: factory number {existing[1]} -> {data['factory_order_no']}. "
                             "Old bills keep the old number.")
            if data.get("quantity") not in ("", None) and Decimal(str(existing[3])) != Decimal(str(data["quantity"] or 0)) and not errs:
                warns.append(f"Stock will be set from {existing[3].normalize():f} to {data['quantity']}.")
        else:
            r["action"] = "CREATE"
            r.pop("product_id", None)
        if sid and data["factory_order_no"]:
            owner = existing_by_fac.get((sid, data["factory_order_no"]))
            if owner and owner[0] != r.get("product_id"):
                errs.append(f"Factory number {data['factory_order_no']} is already mapped to catalogue number "
                            f"{owner[1]} (catalogue {owner[2]}).")
        r["errors"], r["warnings"] = errs, warns

    for (cat, cno), group in seen_cat.items():
        if len(group) > 1:
            rows_txt = ", ".join(str(g["row"]) for g in group)
            for g in group:
                g["errors"].append(f"Catalogue number {cno} appears more than once in this file (rows {rows_txt}).")
    for (_sup, fno), group in seen_fac.items():
        if len(group) > 1:
            rows_txt = ", ".join(str(g["row"]) for g in group)
            for g in group:
                g["errors"].append(f"Factory number {fno} appears more than once in this file (rows {rows_txt}).")

    summary = {"total": len(rows), "errors": sum(1 for r in rows if r["errors"]),
               "warnings": sum(1 for r in rows if r["warnings"] and not r["errors"]),
               "create": sum(1 for r in rows if r["action"] == "CREATE" and not r["errors"]),
               "update": sum(1 for r in rows if r["action"] == "UPDATE" and not r["errors"])}
    return summary


def commit_batch(db: Session, batch: ImportBatch, user: User, skip_errors: bool = False) -> dict:
    rows = batch.rows
    summary = validate_rows(db, rows, batch.mode)
    if summary["errors"] and not skip_errors:
        raise BizError(f"{summary['errors']} rows still have errors. Fix them or choose to skip those rows.")
    good = [r for r in rows if not r["errors"]]
    if not good:
        raise BizError("There are no valid rows to import.")

    cats = {c.name.lower(): c for c in db.execute(select(Category)).scalars()}
    sups = {s.name.lower(): s for s in db.execute(select(Supplier)).scalars()}

    def category_id(name):
        name = (name or "").strip()
        if not name:
            return None
        c = cats.get(name.lower())
        if not c:
            c = Category(name=name)
            db.add(c)
            db.flush()
            cats[name.lower()] = c
        return c.id

    def supplier_id(name):
        s = sups.get(name.lower())
        if not s:
            s = Supplier(name=name)
            db.add(s)
            db.flush()
            sups[name.lower()] = s
        return s.id

    created = updated = 0
    creates = []
    update_ids = [r["product_id"] for r in good if r["action"] == "UPDATE"]
    existing = {}
    for i in range(0, len(update_ids), 900):
        for p in db.execute(select(Product).where(Product.id.in_(update_ids[i:i + 900]))).unique().scalars():
            existing[p.id] = p
    ts = now()
    for r in good:
        x = r["data"]

        def num(f, default="0"):
            v = x.get(f)
            return Decimal(v) if v not in ("", None) else Decimal(default)

        if r["action"] == "CREATE":
            p = Product(catalogue=x["catalogue"], catalogue_no=x["catalogue_no"], factory_order_no=x["factory_order_no"],
                        supplier_id=supplier_id(x["supplier"]), name=x["name"][:200], category_id=category_id(x["category"]),
                        brand=(x.get("brand") or "")[:80], hsn=(x.get("hsn") or "")[:20],
                        description=x.get("description") or "", size=(x.get("size") or "")[:100],
                        finish=(x.get("finish") or "")[:100], unit=(x.get("unit") or "pcs")[:20],
                        sale_price=num("sale_price"), purchase_cost=num("purchase_cost"),
                        tax_rate=num("tax_rate", "18") if x.get("tax_rate") not in ("", None) else Decimal("18"),
                        stock_qty=num("quantity"), reserved_qty=Decimal(0), damaged_qty=Decimal(0),
                        reorder_level=num("reorder_level"), is_active=True, created_at=ts, updated_at=ts)
            creates.append(p)
            created += 1
        else:
            p = existing[r["product_id"]]
            new_sid = supplier_id(x["supplier"])
            if p.factory_order_no != x["factory_order_no"] or p.supplier_id != new_sid:
                db.add(ProductMappingHistory(product_id=p.id, old_catalogue=p.catalogue, old_catalogue_no=p.catalogue_no,
                                             old_factory_order_no=p.factory_order_no, old_supplier_id=p.supplier_id,
                                             new_catalogue=p.catalogue, new_catalogue_no=p.catalogue_no,
                                             new_factory_order_no=x["factory_order_no"], new_supplier_id=new_sid,
                                             reason=f"Import batch {batch.id}", changed_by=user.id))
                p.factory_order_no, p.supplier_id = x["factory_order_no"], new_sid
            p.name = x["name"][:200]
            if x.get("category"):
                p.category_id = category_id(x["category"])
            for f in ("description", "size", "finish", "unit", "brand", "hsn"):
                if x.get(f):
                    setattr(p, f, x[f])
            for f in ("sale_price", "purchase_cost", "tax_rate", "reorder_level"):
                if x.get(f) not in ("", None):
                    setattr(p, f, Decimal(x[f]))
            if x.get("quantity") not in ("", None):
                new_q = Decimal(x["quantity"])
                if new_q < p.reserved_qty:
                    raise BizError(f"Row {r['row']}: stock {new_q} is lower than reserved {p.reserved_qty}.")
                diff = new_q - p.stock_qty
                if diff:
                    p.stock_qty = new_q
                    db.add(StockMovement(product_id=p.id, kind="ADJUST", qty=diff, stock_after=new_q,
                                         reserved_after=p.reserved_qty, ref_type="import", ref_id=batch.id,
                                         note=f"Import batch {batch.id}", user_id=user.id))
            updated += 1
    for i in range(0, len(creates), 2000):
        chunk = creates[i:i + 2000]
        db.add_all(chunk)
        db.flush()
        db.add_all([StockMovement(product_id=p.id, kind="IN", qty=p.stock_qty, stock_after=p.stock_qty,
                                  reserved_after=Decimal(0), ref_type="import", ref_id=batch.id,
                                  note="Opening stock (import)", user_id=user.id, created_at=ts)
                    for p in chunk if p.stock_qty > 0])
        db.flush()
    batch.status = "COMMITTED"
    batch.committed_at = now()
    batch.summary = {**summary, "created": created, "updated": updated, "skipped": len(rows) - len(good)}
    audit(db, user, "import_committed", "import", batch.id, created=created, updated=updated,
          skipped=len(rows) - len(good), filename=batch.filename)
    return batch.summary


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(FIELDS)
    ws.append(["MAIN", "123", "789", "Teak Flush Door", "Luxe", "4418", "Wooden Doors", "Solid core flush door", "7ft x 3ft",
               "Matt", "pcs", "8500", "6200", "18", "25", "5", "Default Factory"])
    ws.append(["MAIN", "117-SM", "RL-117", "Walnut Suede Laminate 1mm", "Luxe", "4823", "Laminates", "", "8x4 ft", "Suede",
               "sheet", "1450", "980", "18", "120", "20", "Default Factory"])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
