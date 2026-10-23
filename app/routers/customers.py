import io
from datetime import datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from ..database import get_db
from ..deps import Page, file_response, opt_date, paged, pdf_response
from ..models import Customer, Invoice, User
from ..pdf import statement_pdf
from ..security import require, require_any
from ..services import (audit, customer_account, customer_dict, customer_totals, next_number, outstanding_by_customer,
                        statement_entries)
from ..types import money

router = APIRouter(prefix="/api/customers", tags=["customers"])
can_view = require_any("customers.view", "billing.create", "payments.record")


class CustomerIn(BaseModel):
    name: str = Field(min_length=1, max_length=150)
    phone: str = Field(default="", max_length=30)
    area: str = Field(default="", max_length=80)
    email: str = ""
    address: str = ""
    gstin: str = Field(default="", max_length=20)
    notes: str = ""
    is_active: bool = True


@router.get("")
def search(q: str = "", with_balance: bool = True, only_outstanding: bool = False, page: Page = Depends(),
           db: Session = Depends(get_db), _: User = Depends(can_view)):
    stmt = select(Customer)
    if q:
        s = q.strip()
        inv_cust = select(Invoice.customer_id).where(Invoice.number == s.upper())
        digits = "".join(ch for ch in s if ch.isdigit())
        conds = [Customer.name.ilike(f"%{s}%"), Customer.area.ilike(f"%{s}%"), Customer.code == s.upper(), Customer.id.in_(inv_cust)]
        if len(digits) >= 3:
            conds.append(Customer.phone.like(f"%{digits}%"))
        stmt = stmt.where(or_(*conds))
    balances = {}
    if with_balance or only_outstanding:
        balances = {r["customer_id"]: r["outstanding"] for r in outstanding_by_customer(db)}
    if only_outstanding:
        stmt = stmt.where(Customer.id.in_([cid for cid, v in balances.items() if v > 0] or [0]))
    total = db.scalar(select(func.count()).select_from(stmt.subquery()))
    rows = db.execute(stmt.order_by(Customer.name).offset(page.offset).limit(page.size)).scalars().all()
    items = []
    for c in rows:
        x = customer_dict(c)
        if balances or with_balance:
            x["outstanding"] = money(balances.get(c.id, 0))
        items.append(x)
    return paged(items, total, page)


@router.post("")
def create(body: CustomerIn, db: Session = Depends(get_db), user: User = Depends(require_any("customers.edit", "billing.create"))):
    phone = body.phone.strip()
    if phone:
        dup = db.execute(select(Customer).where(Customer.phone == phone)).scalars().first()
        if dup:
            raise HTTPException(409, f"Phone {phone} already belongs to {dup.name} ({dup.code}). Open that account instead.")
    c = Customer(code=next_number(db, "C", width=5, yearly=False), **{**body.model_dump(), "name": body.name.strip(),
                                                                      "phone": phone})
    db.add(c)
    db.flush()
    audit(db, user, "customer_created", "customer", c.id, name=c.name, code=c.code)
    db.commit()
    return customer_dict(c)


@router.put("/{customer_id}")
def update(customer_id: int, body: CustomerIn, db: Session = Depends(get_db), user: User = Depends(require("customers.edit"))):
    c = db.get(Customer, customer_id)
    if not c:
        raise HTTPException(404, "Customer not found.")
    phone = body.phone.strip()
    if phone:
        dup = db.execute(select(Customer).where(Customer.phone == phone, Customer.id != c.id)).scalars().first()
        if dup:
            raise HTTPException(409, f"Phone {phone} already belongs to {dup.name} ({dup.code}).")
    for k, v in body.model_dump().items():
        setattr(c, k, v.strip() if isinstance(v, str) else v)
    audit(db, user, "customer_updated", "customer", c.id, name=c.name)
    db.commit()
    return customer_dict(c)


@router.get("/{customer_id}")
def get_one(customer_id: int, db: Session = Depends(get_db), _: User = Depends(can_view)):
    c = db.get(Customer, customer_id)
    if not c:
        raise HTTPException(404, "Customer not found.")
    out = customer_dict(c)
    out["totals"] = {k: money(v) for k, v in customer_totals(db, c.id).items()}
    return out


@router.get("/{customer_id}/account")
def account(customer_id: int, db: Session = Depends(get_db), user: User = Depends(require("customers.view"))):
    return customer_account(db, customer_id, user)


def build_statement(db: Session, customer_id: int, user: User, date_from: Optional[str], date_to: Optional[str]):
    since, upto = opt_date(date_from), opt_date(date_to)
    acc = customer_account(db, customer_id, user)
    ledger = statement_entries(db, customer_id, since, upto)
    totals = {k: money(v) for k, v in customer_totals(db, customer_id, upto=upto).items()}
    return statement_pdf(acc["customer"], totals, ledger, acc["invoices"], since, upto), acc["customer"]["code"]


@router.get("/{customer_id}/statement")
def statement_json(customer_id: int, date_from: Optional[str] = None, date_to: Optional[str] = None,
                   db: Session = Depends(get_db), _: User = Depends(require("customers.view"))):
    return statement_entries(db, customer_id, opt_date(date_from), opt_date(date_to))


@router.get("/{customer_id}/statement.pdf")
def statement(customer_id: int, date_from: Optional[str] = None, date_to: Optional[str] = None,
              db: Session = Depends(get_db), user: User = Depends(require("customers.view"))):
    content, code = build_statement(db, customer_id, user, date_from, date_to)
    return pdf_response(content, f"statement_{code}.pdf")


SOURCE_LABEL = {"MANUAL": "Shop / counter", "ONLINE": "Online", "DELIVERY": "Delivery", "BULK": "Bulk"}
STATUS_LABEL = {"PAID": "Paid", "PARTIAL": "Part paid", "UNPAID": "Unpaid", "CANCELLED": "Cancelled"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def build_customer_xlsx(db: Session, customer_id: int, user: User) -> tuple[bytes, str]:
    """Whole customer account in one Excel file: summary, bills, items, payments."""
    acc = customer_account(db, customer_id, user)
    c, t = acc["customer"], acc["totals"]
    show_factory = user.role != "customer"
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="5A3A22")
    money_fmt = '#,##0.00'

    def sheet(ws, headers, rows, money_cols=(), widths=None):
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append(r)
        for col in money_cols:
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                row[0].number_format = money_fmt
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = (widths or {}).get(h, max(12, len(h) + 4))
        ws.freeze_panes = "A2"

    f = lambda v: float(v)  # noqa: E731
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Customer account"])
    ws["A1"].font = Font(bold=True, size=14)
    for k, v in [("Customer", c["name"]), ("Customer ID", c["code"]), ("Phone", c["phone"]), ("Address", c["address"]),
                 ("GSTIN", c["gstin"]), ("Downloaded on", datetime.now().strftime("%d-%m-%Y %H:%M"))]:
        ws.append([k, v])
    ws.append([])
    ws.append(["Total billed (with GST)", f(t["billed"])])
    ws.append(["Payments received", f(t["paid"])])
    ws.append(["Refunds given", f(t["refunded"])])
    ws.append(["Credit notes", f(t["credited"])])
    ws.append(["Balance to pay" if float(t["outstanding"]) >= 0 else "Advance with us", abs(f(t["outstanding"]))])
    for row in ws.iter_rows(min_row=9, max_row=13, min_col=2, max_col=2):
        row[0].number_format = money_fmt
    ws["A13"].font = ws["B13"].font = Font(bold=True)
    ws.append([])
    ws.append(["By bill type", "Bills", "Total", "Paid", "Due"])
    for cell in ws[ws.max_row]:
        cell.font, cell.fill = head_font, head_fill
    groups = {}
    for i in acc["invoices"]:
        if i["status"] != "ISSUED":
            continue
        g = groups.setdefault(i["source"], [0, 0.0, 0.0, 0.0])
        g[0] += 1
        g[1] += f(i["total"])
        g[2] += f(i["settled"])
        g[3] += f(i["due"])
    for src, (n, tot, paid, due) in groups.items():
        ws.append([SOURCE_LABEL.get(src, src), n, tot, paid, due])
        for col in (3, 4, 5):
            ws.cell(ws.max_row, col).number_format = money_fmt
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 34
    for col in "CDE":
        ws.column_dimensions[col].width = 16

    sheet(wb.create_sheet("Bills"),
          ["Bill no.", "Date", "Bill type", "Subtotal", "Discount", "GST", "Total", "Paid", "Due", "Payment status", "Delivery"],
          [[i["number"], i["invoice_date"], SOURCE_LABEL.get(i["source"], i["source"]), f(i["subtotal"]),
            f(i["discount_total"]), f(i["tax_total"]), f(i["total"]), f(i["settled"]), f(i["due"]),
            STATUS_LABEL.get(i["payment_status"], i["payment_status"]),
            "Delivered" if i["delivery_status"] == "DELIVERED" else "Pending"] for i in reversed(acc["invoices"])],
          money_cols=range(4, 10), widths={"Bill no.": 18, "Payment status": 15})

    item_head = ["Bill no.", "Date", "Bill type", "Catalogue no."] + (["Factory no."] if show_factory else []) + \
        ["Category", "Brand", "Product", "Qty", "Unit", "Rate", "Disc %", "Taxable", "GST %", "GST", "Amount", "Supply"]
    item_rows = []
    for i in reversed(acc["invoices"]):
        if i["status"] != "ISSUED":
            continue
        for ln in i["lines"]:
            item_rows.append([i["number"], i["invoice_date"], SOURCE_LABEL.get(i["source"], i["source"]), ln["catalogue_no"]] +
                             ([ln.get("factory_order_no", "")] if show_factory else []) +
                             [ln.get("category", ""), ln.get("brand", ""), ln["product_name"], f(ln["qty"]), ln["unit"], f(ln["rate"]), f(ln["discount_pct"]),
                              f(ln["taxable"]), f(ln["tax_rate"]), f(ln["tax_amount"]), f(ln["line_total"]),
                              "Factory order" if ln["fulfilment"] == "FACTORY" else "From stock"])
    off = 1 if show_factory else 0
    sheet(wb.create_sheet("Items"), item_head, item_rows,
          money_cols=[11 + off, 13 + off, 15 + off, 16 + off], widths={"Product": 34, "Bill no.": 18, "Category": 16})

    numbers = {i["id"]: i["number"] for i in acc["invoices"]}
    sheet(wb.create_sheet("Payments"), ["Receipt no.", "Date", "Against bill", "Type", "Online / Offline", "Method", "Reference", "Amount", "Status"],
          [[p["number"], p["paid_on"], numbers.get(p["invoice_id"], "Advance"), "Refund" if p["kind"] == "REFUND" else "Payment",
            "Offline" if p["method"] in ("CASH", "CHEQUE") else "Online",
            p["method"], p["reference"], f(p["amount"]), "Void" if p["is_void"] else "OK"] for p in reversed(acc["payments"])],
          money_cols=[8], widths={"Receipt no.": 18, "Against bill": 18, "Reference": 20})

    if acc["credit_notes"]:
        sheet(wb.create_sheet("Credit notes"), ["Credit note no.", "Date", "Against bill", "Reason", "Amount", "Status"],
              [[n["number"], n["note_date"], n["invoice_number"] or "-", n["reason"], f(n["amount"]),
                "Void" if n["is_void"] else "OK"] for n in acc["credit_notes"]], money_cols=[5], widths={"Reason": 34})
    buf = io.BytesIO()
    wb.save(buf)
    safe = "".join(ch if ch.isalnum() else "_" for ch in c["name"])[:40]
    return buf.getvalue(), f"{c['code']}_{safe}_account.xlsx"


@router.get("/{customer_id}/export.xlsx")
def export_account(customer_id: int, db: Session = Depends(get_db), user: User = Depends(require("customers.view"))):
    content, name = build_customer_xlsx(db, customer_id, user)
    return file_response(content, name, XLSX)
