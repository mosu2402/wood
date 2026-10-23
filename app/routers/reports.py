import csv
import io
from datetime import date, timedelta
from decimal import Decimal
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from sqlalchemy import String, cast, func, select
from sqlalchemy.orm import Session

from ..config import settings
from ..database import get_db
from ..deps import file_response, opt_date
from ..models import (ZERO, Category, Customer, Expense, FactoryOrder, Invoice, InvoiceLine, Payment, Product,
                      Purchase, User)
from ..security import require
from ..services import outstanding_by_customer, product_dict
from ..types import d, money, qty

router = APIRouter(prefix="/api", tags=["dashboard & reports"])
perm = require("reports.view")


def _range(date_from, date_to):
    today = date.today()
    since = opt_date(date_from) or today.replace(day=1)
    upto = opt_date(date_to) or today
    if since > upto:
        raise HTTPException(400, "'From' date is after 'To' date.")
    return since, upto


def _sales_totals(db, since, upto, customer_id=None, category_id=None):
    if category_id:
        stmt = (select(func.count(func.distinct(Invoice.id)), func.coalesce(func.sum(InvoiceLine.line_total), 0),
                       func.coalesce(func.sum(InvoiceLine.taxable), 0), func.coalesce(func.sum(InvoiceLine.tax_amount), 0),
                       func.coalesce(func.sum(InvoiceLine.cost_total), 0))
                .join(Invoice, Invoice.id == InvoiceLine.invoice_id).join(Product, Product.id == InvoiceLine.product_id)
                .where(Product.category_id == category_id))
    else:
        stmt = select(func.count(Invoice.id), func.coalesce(func.sum(Invoice.total), 0),
                      func.coalesce(func.sum(Invoice.taxable_total), 0), func.coalesce(func.sum(Invoice.tax_total), 0),
                      func.coalesce(func.sum(Invoice.cost_total), 0))
    stmt = stmt.where(Invoice.status == "ISSUED", Invoice.invoice_date >= since, Invoice.invoice_date <= upto)
    if customer_id:
        stmt = stmt.where(Invoice.customer_id == customer_id)
    n, total, taxable, tax, cost = db.execute(stmt).one()
    return {"invoices": n, "total": d(total), "taxable": d(taxable), "tax": d(tax), "cost": d(cost)}


@router.get("/dashboard")
def dashboard(date_from: Optional[str] = None, date_to: Optional[str] = None, customer_id: Optional[int] = None,
              category_id: Optional[int] = None, db: Session = Depends(get_db), user: User = Depends(perm)):
    since, upto = _range(date_from, date_to)
    show_cost = user.can("costs.view")
    sales = _sales_totals(db, since, upto, customer_id, category_id)

    pay = select(Payment.kind, func.sum(Payment.amount)).where(Payment.is_void.is_(False), Payment.paid_on >= since,
                                                               Payment.paid_on <= upto)
    if customer_id:
        pay = pay.where(Payment.customer_id == customer_id)
    pay_rows = dict(db.execute(pay.group_by(Payment.kind)).all())
    received = d(pay_rows.get("PAYMENT") or 0) - d(pay_rows.get("REFUND") or 0)

    balances = outstanding_by_customer(db, customer_id)
    outstanding_total = sum((b["outstanding"] for b in balances if b["outstanding"] > 0), ZERO)
    advances_total = sum((-b["outstanding"] for b in balances if b["outstanding"] < 0), ZERO)
    top = sorted([b for b in balances if b["outstanding"] > 0], key=lambda b: b["outstanding"], reverse=True)[:10]
    names = {c.id: (c.name, c.code, c.phone) for c in db.execute(
        select(Customer).where(Customer.id.in_([b["customer_id"] for b in top] or [0]))).scalars()}

    low_stmt = select(Product).where(Product.is_active.is_(True), Product.reorder_level > 0,
                                     Product.stock_qty - Product.reserved_qty <= Product.reorder_level)
    if category_id:
        low_stmt = low_stmt.where(Product.category_id == category_id)
    low_count = db.scalar(select(func.count()).select_from(low_stmt.subquery()))
    low = db.execute(low_stmt.order_by((Product.stock_qty - Product.reserved_qty - Product.reorder_level)).limit(10)).unique().scalars().all()

    fo_stmt = select(FactoryOrder.status, func.count()).where(FactoryOrder.status.not_in(["DELIVERED", "CANCELLED"]))
    if customer_id:
        fo_stmt = fo_stmt.where(FactoryOrder.customer_id == customer_id)
    fo_counts = dict(db.execute(fo_stmt.group_by(FactoryOrder.status)).all())

    stock_row = db.execute(select(func.coalesce(func.sum(Product.stock_qty), 0),
                                  func.coalesce(func.sum(Product.reserved_qty), 0)).where(Product.is_active.is_(True))).one()

    trend_stmt = select(Invoice.invoice_date, func.sum(Invoice.total)).where(
        Invoice.status == "ISSUED", Invoice.invoice_date >= max(since, upto - timedelta(days=92)), Invoice.invoice_date <= upto)
    if customer_id:
        trend_stmt = trend_stmt.where(Invoice.customer_id == customer_id)
    trend = [{"date": dt.isoformat(), "total": money(v)} for dt, v in
             db.execute(trend_stmt.group_by(Invoice.invoice_date).order_by(Invoice.invoice_date)).all()]

    out = {
        "range": {"from": since.isoformat(), "to": upto.isoformat()},
        "counts": {"customers": db.scalar(select(func.count(Customer.id))),
                   "products": db.scalar(select(func.count(Product.id)).where(Product.is_active.is_(True))),
                   "invoices": sales["invoices"]},
        "sales_total": money(sales["total"]), "sales_taxable": money(sales["taxable"]), "tax_total": money(sales["tax"]),
        "payments_received": money(received), "refunds": money(pay_rows.get("REFUND") or 0),
        "outstanding_total": money(outstanding_total), "advances_total": money(advances_total),
        "pending_customers": [{"customer_id": b["customer_id"], "name": names.get(b["customer_id"], ("?",))[0],
                               "code": names.get(b["customer_id"], ("", ""))[1],
                               "phone": names.get(b["customer_id"], ("", "", ""))[2],
                               "outstanding": money(b["outstanding"])} for b in top],
        "stock": {"on_hand": qty(stock_row[0]), "reserved": qty(stock_row[1]), "low_stock_count": low_count,
                  "low_stock": [product_dict(p, user, minimal=True) for p in low]},
        "factory_open": fo_counts, "factory_open_total": sum(fo_counts.values()),
        "sales_trend": trend,
    }
    if show_cost:
        pur = select(func.coalesce(func.sum(Purchase.total), 0)).where(
            Purchase.is_void.is_(False), Purchase.purchase_date >= since, Purchase.purchase_date <= upto)
        exp = select(func.coalesce(func.sum(Expense.amount), 0)).where(
            Expense.is_void.is_(False), Expense.expense_date >= since, Expense.expense_date <= upto)
        purchases, expenses = d(db.scalar(pur)), d(db.scalar(exp))
        gross_profit = sales["taxable"] - sales["cost"]
        out.update({"purchases_total": money(purchases), "expenses_total": money(expenses),
                    "cost_of_goods": money(sales["cost"]), "gross_profit": money(gross_profit),
                    "net_profit": money(gross_profit - (ZERO if customer_id or category_id else expenses))})
    return out


def _csv(rows: list[dict], name: str):
    buf = io.StringIO()
    if rows:
        w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)
    return file_response(buf.getvalue().encode("utf-8-sig"), name, "text/csv")


@router.get("/reports/sales")
def sales_report(group: str = Query("date", pattern="^(date|month|customer|product|category)$"),
                 date_from: Optional[str] = None, date_to: Optional[str] = None, customer_id: Optional[int] = None,
                 category_id: Optional[int] = None, format: str = "json", db: Session = Depends(get_db),
                 user: User = Depends(perm)):
    since, upto = _range(date_from, date_to)
    show_cost = user.can("costs.view")
    if group == "date":
        key = [Invoice.invoice_date]
    elif group == "month":
        key = [func.substr(cast(Invoice.invoice_date, String), 1, 7)]
    elif group == "customer":
        key = [Customer.code, Customer.name]
    elif group == "product":
        key = [InvoiceLine.catalogue_no, InvoiceLine.product_name]
    else:
        key = [func.coalesce(Category.name, "Uncategorised")]
    stmt = (select(*key, func.count(func.distinct(Invoice.id)), func.coalesce(func.sum(InvoiceLine.qty), 0),
                   func.sum(InvoiceLine.taxable), func.sum(InvoiceLine.tax_amount), func.sum(InvoiceLine.line_total),
                   func.sum(InvoiceLine.cost_total))
            .select_from(InvoiceLine).join(Invoice, Invoice.id == InvoiceLine.invoice_id)
            .join(Customer, Customer.id == Invoice.customer_id).join(Product, Product.id == InvoiceLine.product_id)
            .outerjoin(Category, Category.id == Product.category_id)
            .where(Invoice.status == "ISSUED", Invoice.invoice_date >= since, Invoice.invoice_date <= upto))
    if customer_id:
        stmt = stmt.where(Invoice.customer_id == customer_id)
    if category_id:
        stmt = stmt.where(Product.category_id == category_id)
    stmt = stmt.group_by(*key).order_by(func.sum(InvoiceLine.line_total).desc() if group not in ("date", "month") else key[0])
    rows = []
    for r in db.execute(stmt.limit(5000)).all():
        k = list(r[:len(key)])
        n, q, taxable, tax, total, cost = r[len(key):]
        row = {"group": " - ".join(str(x) for x in k), "bills": n, "qty": qty(q), "taxable": money(taxable or 0),
               "tax": money(tax or 0), "total": money(total or 0)}
        if show_cost:
            row["cost"] = money(cost or 0)
            row["gross_profit"] = money(d(taxable or 0) - d(cost or 0))
        rows.append(row)
    if format == "csv":
        return _csv(rows, f"sales_by_{group}_{since}_{upto}.csv")
    return {"range": {"from": since.isoformat(), "to": upto.isoformat()}, "group": group, "rows": rows}


@router.get("/reports/outstanding")
def outstanding_report(format: str = "json", min_amount: float = 0.01, db: Session = Depends(get_db),
                       _: User = Depends(perm)):
    balances = {b["customer_id"]: b for b in outstanding_by_customer(db)}
    custs = {c.id: c for c in db.execute(select(Customer).where(Customer.id.in_(list(balances) or [0]))).scalars()}
    rows = []
    for cid, b in sorted(balances.items(), key=lambda kv: kv[1]["outstanding"], reverse=True):
        if b["outstanding"] < Decimal(str(min_amount)):
            continue
        c = custs[cid]
        rows.append({"code": c.code, "name": c.name, "phone": c.phone, "billed": money(b["billed"]),
                     "paid": money(b["paid"]), "refunded": money(b["refunded"]), "credited": money(b["credited"]),
                     "outstanding": money(b["outstanding"]), "customer_id": cid})
    if format == "csv":
        return _csv(rows, "outstanding_balances.csv")
    return {"rows": rows, "total": money(sum((Decimal(r["outstanding"]) for r in rows), ZERO))}


@router.get("/reports/payments")
def payments_report(date_from: Optional[str] = None, date_to: Optional[str] = None, format: str = "json",
                    db: Session = Depends(get_db), _: User = Depends(perm)):
    since, upto = _range(date_from, date_to)
    rows = [{"date": dt.isoformat(), "method": m, "kind": k, "count": n, "amount": money(a)}
            for dt, m, k, n, a in db.execute(
                select(Payment.paid_on, Payment.method, Payment.kind, func.count(), func.sum(Payment.amount))
                .where(Payment.is_void.is_(False), Payment.paid_on >= since, Payment.paid_on <= upto)
                .group_by(Payment.paid_on, Payment.method, Payment.kind).order_by(Payment.paid_on)).all()]
    if format == "csv":
        return _csv(rows, f"payments_{since}_{upto}.csv")
    return {"rows": rows}


@router.get("/reports/stock")
def stock_report(low_only: bool = True, category_id: Optional[int] = None, format: str = "json",
                 db: Session = Depends(get_db), user: User = Depends(perm)):
    stmt = select(Product).where(Product.is_active.is_(True))
    if low_only:
        stmt = stmt.where(Product.reorder_level > 0, Product.stock_qty - Product.reserved_qty <= Product.reorder_level)
    if category_id:
        stmt = stmt.where(Product.category_id == category_id)
    rows = [product_dict(p, user, minimal=True) for p in
            db.execute(stmt.order_by(Product.catalogue_no).limit(5000)).unique().scalars()]
    if format == "csv":
        return _csv([{k: v for k, v in r.items() if k not in ("is_active",)} for r in rows], "stock.csv")
    return {"rows": rows}


# ---------------------------------------------------------------- monthly report in the shop's own sheet format
import calendar  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from pydantic import BaseModel  # noqa: E402

from ..models import AppSetting  # noqa: E402
from ..security import get_current_user, require_owner  # noqa: E402
from ..services import audit, settled_map  # noqa: E402

PAY_LABEL = {"PAID": "Paid", "PARTIAL": "Part paid", "UNPAID": "Pending"}


def _get_setting(db: Session, key: str, default: str = "") -> str:
    row = db.get(AppSetting, key)
    return row.value if row else default


def _month_range(month: Optional[str]):
    m = month or date.today().strftime("%Y-%m")
    try:
        y, mo = int(m[:4]), int(m[5:7])
        start = date(y, mo, 1)
    except (ValueError, IndexError):
        raise HTTPException(400, "Month must look like 2026-10.")
    return m, start, date(y, mo, calendar.monthrange(y, mo)[1])


def monthly_data(db: Session, month: Optional[str]) -> dict:
    m, start, end = _month_range(month)
    invs = db.execute(select(Invoice).where(Invoice.status == "ISSUED", Invoice.invoice_date >= start,
                                            Invoice.invoice_date <= end)
                      .order_by(Invoice.invoice_date, Invoice.id)).unique().scalars().all()
    smap = settled_map(db, [i.id for i in invs])
    rows, areas = [], {}
    gross = gst = net = q = paid = pending = ZERO
    for inv in invs:
        c = inv.customer
        settled = smap[inv.id]
        due = inv.total - settled
        status = "PAID" if settled >= inv.total else ("PARTIAL" if settled > 0 else "UNPAID")
        paid += min(settled, inv.total)
        pending += max(due, ZERO)
        a = areas.setdefault(c.area or "(No area)", {"customers": set(), "bills": 0, "qty": ZERO, "gross": ZERO,
                                                     "gst": ZERO, "net": ZERO, "paid": ZERO, "pending": ZERO})
        a["customers"].add(c.id)
        a["bills"] += 1
        a["paid"] += min(settled, inv.total)
        a["pending"] += max(due, ZERO)
        for ln in inv.lines:
            rows.append({"date": inv.invoice_date.strftime("%d-%m-%Y"), "customer": c.name, "area": c.area or "",
                         "contact": c.phone or "", "category": "Bill" if (inv.tax_mode or "BILL") == "BILL" else "Cash",
                         "brand": ln.brand or "", "design": ln.catalogue_no, "product": ln.product_name,
                         "qty": ln.qty, "unit": ln.unit, "rate": ln.rate, "gross": ln.taxable, "gst": ln.tax_amount,
                         "net": ln.line_total, "bill_no": inv.number, "payment": PAY_LABEL[status],
                         "bill_due": due})
            gross += ln.taxable
            gst += ln.tax_amount
            net += ln.line_total
            q += ln.qty
            a["qty"] += ln.qty
            a["gross"] += ln.taxable
            a["gst"] += ln.tax_amount
            a["net"] += ln.line_total
    target = d(_get_setting(db, "monthly_target", "0") or 0)
    summary = {"month": m, "sales_target": target, "gross": gross, "gst": gst, "net": net, "qty": q,
               "avg_rate": d(gross / q) if q else ZERO, "pending": pending, "paid": paid, "bills": len(invs),
               "achieved_pct": d(net * 100 / target) if target else ZERO,
               "cash_sales": sum((r["net"] for r in rows if r["category"] == "Cash"), ZERO),
               "bill_sales": sum((r["net"] for r in rows if r["category"] == "Bill"), ZERO)}
    area_rows = [{"area": k, "customers": len(v["customers"]), "bills": v["bills"], "qty": v["qty"], "gross": v["gross"],
                  "gst": v["gst"], "net": v["net"], "paid": v["paid"], "pending": v["pending"]}
                 for k, v in sorted(areas.items(), key=lambda kv: kv[1]["net"], reverse=True)]
    return {"summary": summary, "rows": rows, "areas": area_rows}


def _jsonable(v):
    return money(v) if isinstance(v, Decimal) else v


@router.get("/reports/monthly")
def monthly_report(month: Optional[str] = None, db: Session = Depends(get_db), _: User = Depends(perm)):
    data = monthly_data(db, month)
    return {"summary": {k: _jsonable(v) for k, v in data["summary"].items()},
            "rows": [{k: (qty(v) if k == "qty" else _jsonable(v)) for k, v in r.items()} for r in data["rows"]],
            "areas": [{k: (qty(v) if k == "qty" else _jsonable(v)) for k, v in a.items()} for a in data["areas"]]}


@router.get("/reports/monthly.xlsx")
def monthly_excel(month: Optional[str] = None, db: Session = Depends(get_db), _: User = Depends(perm)):
    data = monthly_data(db, month)
    s = data["summary"]
    navy, white = PatternFill("solid", fgColor="1F4E79"), Font(bold=True, color="FFFFFF")
    band = PatternFill("solid", fgColor="DCE6F1")
    rs_fmt = '"₹"#,##0.00'
    wb = Workbook()

    ws = wb.active
    ws.title = "Monthly Sales Tracker"
    head = ["Date", "Dealer / Customer", "Area", "Contact No.", "Category", "Brand", "Design / Code", "Product",
            "Qty (Sheets)", "Unit", "Rate / Sheet", "Gross Sales", "GST", "Net Sales", "Bill No.", "Payment Status"]
    ws.append(head)
    for c in ws[1]:
        c.fill, c.font, c.alignment = navy, white, Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, r in enumerate(data["rows"]):
        ws.append([r["date"], r["customer"], r["area"], r["contact"], r["category"], r["brand"], r["design"], r["product"],
                   float(r["qty"]), r["unit"], float(r["rate"]), float(r["gross"]), float(r["gst"]), float(r["net"]),
                   r["bill_no"], r["payment"]])
        if i % 2 == 0:
            for c in ws[ws.max_row]:
                c.fill = band
    for col in "KLMN":
        for c in ws[col][1:]:
            c.number_format = rs_fmt
    widths = [12, 26, 14, 14, 10, 14, 16, 28, 12, 8, 13, 14, 12, 14, 18, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"
    if data["rows"]:
        ws.auto_filter.ref = f"A1:P{ws.max_row}"

    ws = wb.create_sheet("Monthly Summary")
    ws.append([f"{settings.BUSINESS_NAME} – Monthly Summary"])
    ws.merge_cells("A1:B1")
    ws["A1"].fill, ws["A1"].font = navy, Font(bold=True, color="FFFFFF", size=14)
    items = [("Month", calendar.month_name[int(s["month"][5:7])] + " " + s["month"][:4], None),
             ("Sales Target", float(s["sales_target"]), rs_fmt), ("Total Gross Sales", float(s["gross"]), rs_fmt),
             ("Total GST", float(s["gst"]), rs_fmt), ("Total Net Sales", float(s["net"]), rs_fmt),
             ("Target Achieved %", float(s["achieved_pct"]) / 100, "0.0%"),
             ("Total Quantity (Sheets)", float(s["qty"]), "#,##0.###"), ("Average Rate / Sheet", float(s["avg_rate"]), rs_fmt),
             ("Cash Sales", float(s["cash_sales"]), rs_fmt), ("Bill (GST) Sales", float(s["bill_sales"]), rs_fmt),
             ("Number of Bills", s["bills"], "0"), ("Paid Sales", float(s["paid"]), rs_fmt),
             ("Pending Payments", float(s["pending"]), rs_fmt)]
    for label, val, fmt in items:
        ws.append([label, val])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        if fmt:
            ws.cell(ws.max_row, 2).number_format = fmt
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 28, 20

    ws = wb.create_sheet("Area Wise")
    head = ["Area", "Customers", "Bills", "Qty (Sheets)", "Gross Sales", "GST", "Net Sales", "Paid", "Pending"]
    ws.append(head)
    for c in ws[1]:
        c.fill, c.font = navy, white
    for a in data["areas"]:
        ws.append([a["area"], a["customers"], a["bills"], float(a["qty"]), float(a["gross"]), float(a["gst"]),
                   float(a["net"]), float(a["paid"]), float(a["pending"])])
    if data["areas"]:
        n = ws.max_row
        ws.append(["Total", None, f"=SUM(C2:C{n})", f"=SUM(D2:D{n})", f"=SUM(E2:E{n})", f"=SUM(F2:F{n})",
                   f"=SUM(G2:G{n})", f"=SUM(H2:H{n})", f"=SUM(I2:I{n})"])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for col in "EFGHI":
        for c in ws[col][1:]:
            c.number_format = rs_fmt
    for i, w in enumerate([18, 11, 8, 13, 15, 13, 15, 15, 15], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return file_response(buf.getvalue(), f"Luxe_Monthly_{s['month']}.xlsx",
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


class TargetIn(BaseModel):
    monthly_target: Decimal


@router.get("/settings/shop")
def shop_settings(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    return {"monthly_target": money(d(_get_setting(db, "monthly_target", "0") or 0)),
            "gst_round": settings.GST_ROUND_RUPEE}


@router.put("/settings/shop")
def save_shop_settings(body: TargetIn, db: Session = Depends(get_db), owner: User = Depends(require_owner)):
    if body.monthly_target < 0:
        raise HTTPException(400, "Target cannot be negative.")
    row = db.get(AppSetting, "monthly_target") or AppSetting(key="monthly_target")
    row.value = money(body.monthly_target)
    db.add(row)
    audit(db, owner, "target_changed", "settings", None, monthly_target=row.value)
    db.commit()
    return {"monthly_target": row.value}


# ---------------------------------------------------------------- simple dashboard: orders + payments at a glance
@router.get("/dashboard/simple")
def simple_dashboard(period: str = Query("month", pattern="^(today|month|all)$"), db: Session = Depends(get_db),
                     _: User = Depends(perm)):
    today = date.today()
    since = today if period == "today" else today.replace(day=1) if period == "month" else None
    q = select(Invoice).where(Invoice.status == "ISSUED")
    if since:
        q = q.where(Invoice.invoice_date >= since)
    invs = db.execute(q.order_by(Invoice.invoice_date.desc(), Invoice.id.desc())).unique().scalars().all()
    smap = settled_map(db, [i.id for i in invs])
    total = paid_on_orders = pending = ZERO
    paid_count = 0
    for i in invs:
        st = smap[i.id]
        total += i.total
        paid_on_orders += min(st, i.total)
        pending += max(i.total - st, ZERO)
        paid_count += 1 if st >= i.total else 0
    pq = select(Payment.kind, func.coalesce(func.sum(Payment.amount), 0)).where(Payment.is_void.is_(False))
    if since:
        pq = pq.where(Payment.paid_on >= since)
    pr = dict(db.execute(pq.group_by(Payment.kind)).all())
    received = d(pr.get("PAYMENT") or 0) - d(pr.get("REFUND") or 0)
    not_delivered = [i for i in invs if i.delivery_status != "DELIVERED"]
    balances = sorted([b for b in outstanding_by_customer(db) if b["outstanding"] > 0], key=lambda b: b["outstanding"], reverse=True)
    names = {c.id: c for c in db.execute(select(Customer).where(Customer.id.in_([b["customer_id"] for b in balances[:15]] or [0]))).scalars()}

    def row(i):
        st = smap[i.id]
        return {"id": i.id, "number": i.number, "date": i.invoice_date.isoformat(), "customer_id": i.customer_id,
                "customer": i.customer.name, "area": i.customer.area or "", "total": money(i.total), "paid": money(min(st, i.total)),
                "due": money(max(i.total - st, ZERO)), "payment_status": "PAID" if st >= i.total else ("PARTIAL" if st > 0 else "UNPAID"),
                "delivered": i.delivery_status == "DELIVERED", "tax_mode": i.tax_mode or "BILL"}

    return {
        "period": period,
        "orders": len(invs), "orders_pending": len(not_delivered), "orders_delivered": len(invs) - len(not_delivered),
        "orders_paid": paid_count, "orders_unpaid": len(invs) - paid_count,
        "sales_total": money(total), "payment_received": money(received), "payment_pending": money(pending),
        "customers_owe_total": money(sum((b["outstanding"] for b in balances), ZERO)),
        "to_deliver": [row(i) for i in sorted(not_delivered, key=lambda x: (x.invoice_date, x.id))[:15]],
        "to_collect": [{"customer_id": b["customer_id"], "name": names[b["customer_id"]].name, "phone": names[b["customer_id"]].phone,
                        "area": names[b["customer_id"]].area or "", "pending": money(b["outstanding"])}
                       for b in balances[:15] if b["customer_id"] in names],
        "recent": [row(i) for i in invs[:10]],
    }
